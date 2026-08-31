from openpyxl import load_workbook
from typing import List, Dict, Any

def extract_camels_data(excel_file: str) -> List[Dict[str, Any]]:
    """
    Đọc toàn bộ cấu trúc file Excel CAMELS bao gồm cả Nhóm chỉ tiêu (Header level 1, 2)
    và Các chỉ tiêu tính toán & dữ liệu đầu vào (Level 3, 4, R-191, R-192) kèm phân cấp level.
    """
    wb = load_workbook(excel_file, data_only=True)
    ws = wb.active

    HEADER_ROW = 7

    headers = [cell.value for cell in ws[HEADER_ROW]]
    header_index = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}

    ma_idx = header_index.get("Mã dòng nguồn", 0)
    stt_idx = header_index.get("STT", 1)
    chi_tieu_idx = header_index.get("Chỉ tiêu", 2)
    so_lieu_idx = header_index.get("Số liệu", 3)

    result = []

    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if all(v is None for v in row):
            continue

        code = row[ma_idx]
        if not code or not str(code).strip().startswith("R-"):
            continue

        stt = row[stt_idx]
        chi_tieu = row[chi_tieu_idx]
        so_lieu = row[so_lieu_idx]

        stt_str = str(stt).strip() if stt is not None else ""
        chi_tieu_str = str(chi_tieu).strip() if chi_tieu is not None else ""
        code_str = str(code).strip()

        stt_parts = [p for p in stt_str.split(".") if p]
        level = len(stt_parts) if stt_parts else 1

        # Cấp 1 (VD: 1, 2...) và Cấp 2 (VD: 1.1, 1.2...) là các Nhóm tiêu chí
        # Riêng R-191, R-192 và Cấp 3, 4 trở đi là các Chỉ tiêu thành phần (is_header = False)
        is_header = (level == 1) or (level == 2 and code_str not in ["R-191", "R-192"])

        result.append({
            "ma_dong_nguon": code_str,
            "stt": stt_str,
            "chi_tieu": chi_tieu_str,
            "so_lieu": so_lieu,
            "cap_do": level,
            "is_header": is_header,
        })

    return result